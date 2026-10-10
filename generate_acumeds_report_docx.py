from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(__file__).resolve().parent
MARKDOWN_PATH = ROOT / 'AcuMeds_Project_Report.md'
OUTPUT_PATH = ROOT / 'AcuMeds_Project_Report.docx'


def clean_text(text: str) -> str:
    text = text.replace('**', '')
    text = text.replace('`', '')
    return text.strip()


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_default_styles(document: Document) -> None:
    normal = document.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)

    for style_name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = document.styles[style_name]
        style.font.name = 'Times New Roman'

    document.styles['Title'].font.size = Pt(22)
    document.styles['Title'].font.bold = True
    document.styles['Heading 1'].font.size = Pt(16)
    document.styles['Heading 1'].font.bold = True
    document.styles['Heading 2'].font.size = Pt(14)
    document.styles['Heading 2'].font.bold = True
    document.styles['Heading 3'].font.size = Pt(12)
    document.styles['Heading 3'].font.bold = True


def set_margins(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def set_page_numbering(section, start_value: int = 0) -> None:
    pg_num_type = section._sectPr.find(qn('w:pgNumType'))
    if pg_num_type is None:
        pg_num_type = OxmlElement('w:pgNumType')
        section._sectPr.append(pg_num_type)
    pg_num_type.set(qn('w:start'), str(start_value))
    pg_num_type.set(qn('w:fmt'), 'decimal')


def add_page_number_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')

    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '

    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')

    text = OxmlElement('w:t')
    text.text = '1'

    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')

    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')

    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '

    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')

    text = OxmlElement('w:t')
    text.text = 'Right-click and update the table of contents in Word.'

    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')

    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def add_centered_paragraph(document: Document, text: str, size: int, bold: bool = False, italic: bool = False, space_after: int = 0):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_after)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return paragraph


def add_bullets(document: Document, items):
    for item in items:
        paragraph = document.add_paragraph(style='List Bullet')
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(clean_text(item))
        set_run_font(run)


def add_numbered_list(document: Document, items):
    for item in items:
        paragraph = document.add_paragraph(style='List Number')
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(clean_text(item))
        set_run_font(run)


def add_plain_paragraph(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(clean_text(text))
    set_run_font(run)


def add_table(document: Document, table_lines):
    rows = []
    for line in table_lines:
        cells = [clean_text(cell) for cell in line.strip().strip('|').split('|')]
        rows.append(cells)

    if not rows:
        return

    header = rows[0]
    data_rows = rows[2:] if len(rows) >= 2 and all(set(part.strip()) <= {'-', ':'} for part in rows[1]) else rows[1:]
    table = document.add_table(rows=1, cols=len(header))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for index, value in enumerate(header):
        hdr_cells[index].text = value

    for row in data_rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if index < len(cells):
                cells[index].text = value


def add_title_page(document: Document):
    document.add_paragraph('')
    document.add_paragraph('')
    add_centered_paragraph(document, 'AcuMeds', 24, bold=True, space_after=6)
    add_centered_paragraph(document, 'Integrated Online Pharmacy Management System', 16, bold=True, space_after=2)
    add_centered_paragraph(document, 'with AI-Based Recommendation and Multi-Portal Workflow', 16, bold=True, space_after=12)
    add_centered_paragraph(document, 'Project Report', 15, bold=True, space_after=12)
    add_centered_paragraph(document, 'Submitted in partial fulfillment of the requirements for the award of [Degree Name]', 12, space_after=10)
    add_centered_paragraph(document, 'Submitted by:', 12, bold=True, space_after=2)
    add_centered_paragraph(document, '[Your Name]', 12, space_after=2)
    add_centered_paragraph(document, '[Roll Number]', 12, space_after=2)
    add_centered_paragraph(document, '[Department / Program]', 12, space_after=10)
    add_centered_paragraph(document, 'Under the guidance of:', 12, bold=True, space_after=2)
    add_centered_paragraph(document, '[Guide Name]', 12, space_after=2)
    add_centered_paragraph(document, '[Institution Name]', 12, space_after=2)
    add_centered_paragraph(document, '[Academic Session]', 12, space_after=2)


def parse_markdown_body(document: Document, markdown_text: str) -> None:
    lines = markdown_text.splitlines()
    start_index = None
    for index, line in enumerate(lines):
        if line.startswith('## Declaration'):
            start_index = index
            break

    if start_index is None:
        raise RuntimeError('Could not find the Declaration section in the markdown report.')

    parsing_index = False
    paragraph_buffer = []
    bullet_buffer = []
    numbered_buffer = []
    table_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            add_plain_paragraph(document, ' '.join(paragraph_buffer))
            paragraph_buffer = []

    def flush_bullets():
        nonlocal bullet_buffer
        if bullet_buffer:
            add_bullets(document, bullet_buffer)
            bullet_buffer = []

    def flush_numbered():
        nonlocal numbered_buffer
        if numbered_buffer:
            add_numbered_list(document, numbered_buffer)
            numbered_buffer = []

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            add_table(document, table_buffer)
            table_buffer = []

    def flush_all():
        flush_paragraph()
        flush_bullets()
        flush_numbered()
        flush_table()

    for raw_line in lines[start_index:]:
        line = raw_line.rstrip()

        if parsing_index:
            if line.startswith('## List of Tables'):
                parsing_index = False
            else:
                continue

        if line.strip() == '\\newpage' or line.strip() == '---':
            flush_all()
            document.add_page_break()
            continue

        if line.startswith('## Index'):
            flush_all()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run('Index')
            set_run_font(run, size=16, bold=True)

            toc_paragraph = document.add_paragraph()
            toc_paragraph.paragraph_format.space_after = Pt(0)
            add_toc_field(toc_paragraph)
            document.add_paragraph('Update the table of contents in Word to refresh page numbers.')
            document.paragraphs[-1].paragraph_format.space_after = Pt(0)
            document.paragraphs[-1].paragraph_format.line_spacing = 1.2
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            parsing_index = True
            continue

        heading_match = re.match(r'^(#{2,4})\s+(.+)$', line)
        if heading_match:
            flush_all()
            level = len(heading_match.group(1)) - 1
            title = clean_text(heading_match.group(2))
            document.add_heading(title, level=level)
            continue

        if not line.strip():
            flush_all()
            continue

        if re.match(r'^\s*-\s+', line):
            flush_paragraph()
            flush_numbered()
            bullet_buffer.append(line.replace('- ', '', 1).strip())
            continue

        if re.match(r'^\s*\d+\.\s+', line):
            flush_paragraph()
            flush_bullets()
            numbered_buffer.append(re.sub(r'^\s*\d+\.\s+', '', line).strip())
            continue

        if line.startswith('|'):
            flush_paragraph()
            flush_bullets()
            flush_numbered()
            table_buffer.append(line)
            continue

        flush_bullets()
        flush_numbered()
        flush_table()
        paragraph_buffer.append(clean_text(line))

    flush_all()


def build_document() -> Document:
    document = Document()
    set_default_styles(document)
    set_margins(document)

    section = document.sections[0]
    section.different_first_page_header_footer = True
    set_page_numbering(section, start_value=0)

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(0)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    footer_run = footer_paragraph.add_run('Page ')
    set_run_font(footer_run, size=10)
    add_page_number_field(footer_paragraph)

    add_title_page(document)
    document.add_page_break()

    markdown_text = MARKDOWN_PATH.read_text(encoding='utf-8')
    parse_markdown_body(document, markdown_text)

    return document


def main():
    document = build_document()
    document.save(OUTPUT_PATH)
    print(f'Created {OUTPUT_PATH}')


if __name__ == '__main__':
    main()