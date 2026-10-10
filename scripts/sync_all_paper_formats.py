"""
Synchronizes the expanded 9-page IEEE research paper into DOCX and Markdown formats.
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path("D:/CODING/AI MEDICAL")
HTML_P = ROOT / "AcuMeds_AI_Medical_IEEE_Paper.html"
DOCX_P = ROOT / "AcuMeds_AI_Medical_IEEE_Paper.docx"
MD_P = ROOT / "AcuMeds_AI_Medical_IEEE_Paper.md"

def sync():
    with open(HTML_P, "r", encoding="utf-8") as f:
        html = f.read()

    # 1. Build DOCX
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_head.add_run("AcuMeds: Integrated Online Pharmacy Management System with AI-Based Recommendation\n\n")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(18)
    r_t.font.bold = True

    # Find sections
    sections = re.findall(r'<h2 class="section-heading">(.*?)</h2>(.*?)(?=(?:<h2 class="section-heading">)|(?:</div>\s*</body>))', html, re.DOTALL)
    
    md_lines = [
        "# AcuMeds: Integrated Online Pharmacy Management System with AI-Based Recommendation\n",
        "---\n"
    ]

    for sec_title, sec_body in sections:
        clean_sec_title = sec_title.strip()
        
        # Add to DOCX
        p_s = doc.add_paragraph()
        r_sec = p_s.add_run(clean_sec_title)
        r_sec.font.name = "Times New Roman"
        r_sec.font.size = Pt(12)
        r_sec.font.bold = True
        p_s.paragraph_format.space_before = Pt(12)
        p_s.paragraph_format.space_after = Pt(4)

        # Add to Markdown
        md_lines.append(f"\n## {clean_sec_title}\n")

        # Subsections
        subsections = re.findall(r'<h3 class="subsection-heading">(.*?)</h3>', sec_body)
        paragraphs = re.findall(r'<p.*?>(.*?)</p>', sec_body, re.DOTALL)

        for p in paragraphs:
            clean_p = re.sub(r'<.*?>', '', p).strip()
            if clean_p and not clean_p.startswith('Abstract'):
                p_para = doc.add_paragraph()
                r_p = p_para.add_run(clean_p)
                r_p.font.name = "Times New Roman"
                r_p.font.size = Pt(10)
                p_para.paragraph_format.space_after = Pt(4)

                md_lines.append(f"{clean_p}\n")

    doc.save(DOCX_P)
    with open(MD_P, "w", encoding="utf-8") as f:
        f.writelines(md_lines)

    print("SUCCESS: Synchronized DOCX and Markdown files!")

if __name__ == "__main__":
    sync()
