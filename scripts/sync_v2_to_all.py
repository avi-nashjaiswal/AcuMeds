"""
Syncs the complete publication-grade IEEE research paper into DOCX and Markdown.
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path("D:/CODING/AI MEDICAL")
HTML_P = ROOT / "AcuMeds_AI_Medical_IEEE_Paper.html"
DOCX_P = ROOT / "AcuMeds_AI_Medical_IEEE_Paper.docx"
MD_P = ROOT / "AcuMeds_AI_Medical_IEEE_Paper.md"

def sync():
    with open(HTML_P, "r", encoding="utf-8") as f:
        html = f.read()

    # 1. Build DOCX
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_head.add_run("AcuMeds: An AI-Powered Smart E-Pharmacy and Doctor Consultation System with Explainable Medicine Recommendation\n\n")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(18)
    r_t.font.bold = True

    # Abstract
    abs_m = re.search(r'<span class="abstract-title">Abstract</span>—(.*?)</p>', html, re.DOTALL)
    abs_text = re.sub(r'<.*?>', '', abs_m.group(1)).strip() if abs_m else ''
    p_abs = doc.add_paragraph()
    r1 = p_abs.add_run("Abstract—")
    r1.font.bold = True
    r1.font.italic = True
    p_abs.add_run(abs_text + "\n")

    p_kw = doc.add_paragraph()
    r2 = p_kw.add_run("Keywords—")
    r2.font.bold = True
    r2.font.italic = True
    p_kw.add_run("Explainable Artificial Intelligence (XAI), Clinical Decision Support Systems (CDSS), Smart E-Pharmacy, Telemedicine, WebRTC Video Streaming, SHAP, LIME, Natural Language Processing, UPI Fraud Mitigation, Digital Prescriptions.\n")

    # Add Sections
    sections = re.findall(r'<h2 class="section-heading">(.*?)</h2>(.*?)(?=(?:<h2 class="section-heading">)|(?:</div>\s*</body>))', html, re.DOTALL)
    
    md_lines = [
        "# AcuMeds: An AI-Powered Smart E-Pharmacy and Doctor Consultation System with Explainable Medicine Recommendation\n\n",
        "---\n\n",
        f"**Abstract**—{abs_text}\n\n",
        "**Keywords:** *Explainable Artificial Intelligence (XAI), Clinical Decision Support Systems (CDSS), Smart E-Pharmacy, Telemedicine, WebRTC Video Streaming, SHAP, LIME, Natural Language Processing, UPI Fraud Mitigation, Digital Prescriptions.*\n\n",
        "---\n"
    ]

    for sec_title, sec_body in sections:
        clean_title = sec_title.strip()
        p_sec = doc.add_paragraph()
        r_s = p_sec.add_run(clean_title)
        r_s.font.name = "Times New Roman"
        r_s.font.size = Pt(12)
        r_s.font.bold = True
        p_sec.paragraph_format.space_before = Pt(12)
        p_sec.paragraph_format.space_after = Pt(4)

        md_lines.append(f"\n## {clean_title}\n\n")

        # Subsections and Paragraphs
        parts = re.findall(r'(<h3 class="subsection-heading">.*?</h3>|<p.*?>.*?</p>|<div class="algorithm-box">.*?</div>)', sec_body, re.DOTALL)
        for part in parts:
            if part.startswith('<h3'):
                sub_t = re.sub(r'<.*?>', '', part).strip()
                p_sub = doc.add_paragraph()
                r_sub = p_sub.add_run(sub_t)
                r_sub.font.name = "Times New Roman"
                r_sub.font.size = Pt(11)
                r_sub.font.bold = True
                r_sub.font.italic = True
                md_lines.append(f"\n### {sub_t}\n\n")
            elif part.startswith('<div class="algorithm-box">'):
                code_t = re.sub(r'<.*?>', '', part).strip()
                p_c = doc.add_paragraph()
                r_c = p_c.add_run(code_t)
                r_c.font.name = "Courier New"
                r_c.font.size = Pt(8.5)
                md_lines.append(f"```text\n{code_t}\n```\n\n")
            else:
                p_clean = re.sub(r'<.*?>', '', part).strip()
                if p_clean and not p_clean.startswith('Abstract') and not p_clean.startswith('Index Terms') and not p_clean.startswith('Keywords'):
                    p_p = doc.add_paragraph()
                    r_p = p_p.add_run(p_clean)
                    r_p.font.name = "Times New Roman"
                    r_p.font.size = Pt(10)
                    p_p.paragraph_format.space_after = Pt(4)
                    md_lines.append(f"{p_clean}\n\n")

    # References
    refs = re.findall(r'<div class="reference-item">(.*?)</div>', html, re.DOTALL)
    if refs:
        p_r = doc.add_paragraph()
        r_rh = p_r.add_run("References")
        r_rh.font.name = "Times New Roman"
        r_rh.font.size = Pt(12)
        r_rh.font.bold = True
        md_lines.append("\n## References\n\n")
        for ref in refs:
            ref_clean = re.sub(r'<.*?>', '', ref).strip()
            p_ref = doc.add_paragraph()
            r_rf = p_ref.add_run(ref_clean)
            r_rf.font.name = "Times New Roman"
            r_rf.font.size = Pt(9)
            md_lines.append(f"- {ref_clean}\n")

    doc.save(DOCX_P)
    with open(MD_P, "w", encoding="utf-8") as f:
        f.writelines(md_lines)

    print("SUCCESS: Synchronized DOCX and Markdown formats!")

if __name__ == "__main__":
    sync()
