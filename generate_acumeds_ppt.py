from pathlib import Path
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor

ROOT = Path(r"D:\CODING\AI MEDICAL")
HOME_PNG = ROOT / "AcuMeds_home_preview.png"
AI_PNG = ROOT / "AcuMeds_recommendation_preview.png"
GALLERY_PNG = ROOT / "AcuMeds_medicine_gallery.png"
INTRODUCTION_PNG = ROOT / "AcuMeds_introduction_image.png"
PROBLEM_PNG = ROOT / "AcuMeds_problem_statement_image.png"
PROPOSED_PNG = ROOT / "AcuMeds_proposed_solution_image.png"
ARCHITECTURE_PNG = ROOT / "AcuMeds_system_architecture_image.png"
AI_MODULE_PNG = ROOT / "AcuMeds_ai_recommendation_module_image.png"
USER_ROLES_PNG = ROOT / "AcuMeds_user_roles_image.png"
WORKFLOW_PNG = ROOT / "AcuMeds_workflow_image.png"
TECHNOLOGIES_PNG = ROOT / "AcuMeds_technologies_image.png"
FEATURES_PNG = ROOT / "AcuMeds_features_image.png"
RESULT_PNG = ROOT / "AcuMeds_result_image.png"
BENEFITS_PNG = ROOT / "AcuMeds_benefits_image.png"
DIFFERENCE_PNG = ROOT / "AcuMeds_difference_image.png"
SCOPE_PNG = ROOT / "AcuMeds_scope_image.png"
CONCLUSION_PNG = ROOT / "AcuMeds_conclusion_image.png"
OUT = ROOT / "AcuMeds_Project_Presentation_Interactive.pptx"

SECTION_IMAGE_PATHS = {
    "Introduction": INTRODUCTION_PNG,
    "Problem Statement": PROBLEM_PNG,
    "Proposed Solution": PROPOSED_PNG,
    "System Architecture": ARCHITECTURE_PNG,
    "AI Recommendation Module": AI_MODULE_PNG,
    "User Roles": USER_ROLES_PNG,
    "Workflow": WORKFLOW_PNG,
    "Technologies": TECHNOLOGIES_PNG,
    "Features": FEATURES_PNG,
    "Result": RESULT_PNG,
    "Benefits": BENEFITS_PNG,
    "How Different From Other Systems": DIFFERENCE_PNG,
    "Scope": SCOPE_PNG,
    "Conclusion": CONCLUSION_PNG,
}

MEDICINE_SOURCE_IMAGES = [
    (ROOT / "shopkeeper_system" / "static" / "uploads" / "medicines" / "medicine_3_dolo-650-mg-paracetamol-tablet.jpeg", "Paracetamol", (239, 133, 53)),
    (ROOT / "shopkeeper_system" / "static" / "uploads" / "medicines" / "medicine_5_download_2.jpg", "Medicine Preview", (38, 122, 183)),
    (ROOT / "shopkeeper_system" / "static" / "uploads" / "medicines" / "medicine_2_download_8.jpg", "Medicine Preview", (71, 171, 151)),
    (ROOT / "shopkeeper_system" / "static" / "uploads" / "medicines" / "medicine_19_download.jpg", "Medicine Preview", (24, 107, 118)),
    (HOME_PNG, "Homepage UI", (38, 122, 183)),
    (AI_PNG, "AI Module", (71, 171, 151)),
]

NAVY = RGBColor(10, 30, 46)
TEAL = RGBColor(24, 107, 118)
BLUE = RGBColor(38, 122, 183)
ORANGE = RGBColor(239, 133, 53)
MINT = RGBColor(71, 171, 151)
BG = RGBColor(245, 248, 252)
CARD = RGBColor(255, 255, 255)
LINE = RGBColor(224, 231, 239)
TEXT = RGBColor(37, 49, 58)
MUTED = RGBColor(92, 106, 122)
WHITE = RGBColor(255, 255, 255)
FONT = "Aptos"


def font_path(candidates):
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return str(path)
    return None


REG_FONT = font_path([
    r"C:\Windows\Fonts\segoeui.ttf",
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\calibri.ttf",
])
BOLD_FONT = font_path([
    r"C:\Windows\Fonts\segoeuib.ttf",
    r"C:\Windows\Fonts\arialbd.ttf",
    r"C:\Windows\Fonts\calibrib.ttf",
])


def fnt(size, bold=False):
    path = BOLD_FONT if bold else REG_FONT
    if path:
        try:
            return ImageFont.truetype(path, size=size)
        except Exception:
            pass
    return ImageFont.load_default()


# ---------- Website preview images ----------
def browser_chrome(draw, title):
    draw.rectangle((0, 0, 1600, 72), fill=(28, 36, 51))
    for i, c in enumerate([(255, 95, 86), (255, 189, 46), (39, 201, 63)]):
        draw.ellipse((20 + i * 24, 22, 36 + i * 24, 38), fill=c)
    draw.rounded_rectangle((120, 15, 1480, 52), radius=18, fill=(46, 56, 74))
    draw.text((145, 22), title, font=fnt(18), fill=(235, 240, 247))


def logo(draw, x, y):
    draw.ellipse((x, y, x + 42, y + 42), fill=(28, 107, 118))
    draw.rectangle((x + 18, y + 9, x + 24, y + 33), fill=WHITE)
    draw.rectangle((x + 9, y + 18, x + 33, y + 24), fill=WHITE)
    draw.text((x + 54, y + 8), "AcuMeds", font=fnt(24, True), fill=(18, 38, 56))


def pill(draw, box, fill):
    x0, y0, x1, y1 = box
    draw.rounded_rectangle(box, radius=(y1 - y0) // 2, fill=fill)
    mid = (x0 + x1) // 2
    draw.line((mid, y0 + 10, mid, y1 - 10), fill=WHITE, width=4)


def shadow_card(draw, box, fill=WHITE, outline=(220, 225, 232), shadow=(0, 0, 0, 28), offset=(5, 6), radius=20):
    x0, y0, x1, y1 = box
    sx, sy = offset
    draw.rounded_rectangle((x0 + sx, y0 + sy, x1 + sx, y1 + sy), radius=radius, fill=shadow)
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=2)


def home_preview(path):
    img = Image.new("RGBA", (1600, 1000), (245, 248, 252, 255))
    d = ImageDraw.Draw(img, "RGBA")
    browser_chrome(d, "AcuMeds - Home")
    d.rectangle((0, 72, 1600, 1000), fill=(245, 248, 252))
    d.rectangle((0, 72, 1600, 160), fill=WHITE)
    logo(d, 52, 92)
    for i, label in enumerate(["Home", "Medicines", "AI Check", "Orders", "Profile"]):
        d.text((910 + i * 110, 108), label, font=fnt(18), fill=(74, 90, 108))
    d.rounded_rectangle((1460, 98, 1548, 142), radius=16, fill=(38, 122, 183))
    d.text((1481, 108), "Login", font=fnt(18, True), fill=WHITE)

    d.rounded_rectangle((40, 182, 1560, 484), radius=30, fill=(228, 239, 255))
    d.rounded_rectangle((60, 202, 760, 462), radius=26, fill=(240, 247, 255))
    d.text((90, 240), "AI-Powered Online Pharmacy", font=fnt(48, True), fill=(14, 34, 52))
    d.text((90, 315), "Browse medicines, check symptoms,\nand place orders in one place.", font=fnt(23), fill=(89, 104, 122))
    d.rounded_rectangle((90, 388, 245, 440), radius=18, fill=(38, 122, 183))
    d.text((125, 401), "Shop Now", font=fnt(20, True), fill=WHITE)
    d.rounded_rectangle((260, 388, 420, 440), radius=18, fill=WHITE, outline=(38, 122, 183), width=2)
    d.text((294, 401), "AI Check", font=fnt(20, True), fill=(38, 122, 183))

    d.rounded_rectangle((980, 202, 1495, 442), radius=24, fill=WHITE, outline=(223, 230, 240), width=2)
    d.text((1020, 228), "Quick Symptom Check", font=fnt(24, True), fill=(14, 34, 52))
    d.text((1020, 266), "Type: fever, cold, cough, headache", font=fnt(17), fill=(108, 123, 142))
    bars = [("Fever", 86, (239, 133, 53)), ("Cold", 74, (24, 107, 118)), ("Cough", 67, (71, 171, 151))]
    for idx, (label, pct, c) in enumerate(bars):
        y = 316 + idx * 48
        d.text((1020, y), label, font=fnt(18, True), fill=(54, 68, 86))
        d.rounded_rectangle((1100, y + 6, 1450, y + 24), radius=10, fill=(233, 238, 245))
        d.rounded_rectangle((1100, y + 6, 1100 + int(3.5 * pct), y + 24), radius=10, fill=c)
    d.rounded_rectangle((1020, 410, 1140, 440), radius=14, fill=(71, 171, 151))
    d.text((1043, 417), "AI Match", font=fnt(15, True), fill=WHITE)
    d.rounded_rectangle((1160, 410, 1360, 440), radius=14, fill=(246, 250, 254), outline=(216, 225, 235), width=1)
    d.text((1180, 417), "Confidence 82%", font=fnt(15, True), fill=(38, 122, 183))

    cards = [
        ("Secure Login", "Passwords + roles", (239, 133, 53)),
        ("Delivery Tracking", "Status updates", (38, 122, 183)),
        ("AI Suggestions", "Symptom matching", (71, 171, 151)),
    ]
    for i, (title, sub, c) in enumerate(cards):
        x0 = 60 + i * 500
        shadow_card(d, (x0, 520, x0 + 460, 680), fill=WHITE)
        d.ellipse((x0 + 28, 548, x0 + 94, 614), fill=c)
        d.text((x0 + 120, 552), title, font=fnt(24, True), fill=(20, 36, 55))
        d.text((x0 + 120, 592), sub, font=fnt(17), fill=(94, 108, 124))
        d.text((x0 + 28, 632), "Reliable and simple for daily use.", font=fnt(16), fill=(94, 108, 124))

    d.text((60, 720), "Featured Medicines", font=fnt(28, True), fill=(18, 38, 56))
    meds = [
        ("Paracetamol", "Pain relief", "$3.50", (255, 243, 235), (239, 133, 53), True),
        ("Cough Syrup", "Cold & cough", "$4.25", (235, 247, 248), (24, 107, 118), False),
        ("Antacid", "Stomach relief", "$3.95", (240, 248, 243), (71, 171, 151), False),
    ]
    for i, (name, cat, price, bg, accent, rx) in enumerate(meds):
        x0 = 60 + i * 490
        shadow_card(d, (x0, 770, x0 + 430, 956), fill=WHITE)
        d.rounded_rectangle((x0, 770, x0 + 430, 856), radius=22, fill=bg)
        pill(d, (x0 + 26, 792, x0 + 128, 842), fill=accent)
        d.ellipse((x0 + 152, 792, x0 + 186, 826), fill=accent)
        d.ellipse((x0 + 192, 806, x0 + 216, 830), fill=accent)
        d.text((x0 + 236, 790), name, font=fnt(23, True), fill=(18, 38, 56))
        d.text((x0 + 236, 826), cat, font=fnt(16), fill=(94, 108, 124))
        d.text((x0 + 26, 874), price, font=fnt(22, True), fill=(31, 146, 84))
        if rx:
            d.rounded_rectangle((x0 + 310, 868, x0 + 398, 900), radius=12, fill=(255, 193, 7))
            d.text((x0 + 324, 875), "Rx", font=fnt(16, True), fill=(33, 37, 41))
        d.rounded_rectangle((x0 + 292, 914, x0 + 404, 948), radius=12, fill=accent)
        d.text((x0 + 309, 920), "Add to cart", font=fnt(16, True), fill=WHITE)

    d.rectangle((0, 962, 1600, 1000), fill=(29, 38, 54))
    d.text((60, 972), "AcuMeds • Homepage preview", font=fnt(16), fill=(242, 246, 250))
    img.save(path)


def ai_preview(path):
    img = Image.new("RGBA", (1600, 1000), (245, 248, 252, 255))
    d = ImageDraw.Draw(img, "RGBA")
    browser_chrome(d, "AcuMeds - AI Recommendation")
    d.rectangle((0, 72, 1600, 1000), fill=(245, 248, 252))
    d.rectangle((0, 72, 1600, 160), fill=WHITE)
    logo(d, 52, 92)
    d.text((1060, 108), "Recommendations", font=fnt(18), fill=(74, 90, 108))
    d.text((1240, 108), "History", font=fnt(18), fill=(74, 90, 108))
    d.rounded_rectangle((1400, 98, 1510, 142), radius=16, fill=(71, 171, 151))
    d.text((1426, 108), "Results", font=fnt(18, True), fill=WHITE)

    d.rounded_rectangle((40, 182, 1560, 940), radius=30, fill=WHITE, outline=(223, 230, 240), width=2)
    d.text((70, 214), "AI Recommendation Module", font=fnt(42, True), fill=(18, 38, 56))
    d.text((70, 270), "Enter symptoms and get ranked medicine suggestions based on live inventory.", font=fnt(21), fill=(96, 110, 126))

    shadow_card(d, (70, 330, 680, 860), fill=(248, 250, 253, 255))
    d.text((100, 356), "Symptom Input", font=fnt(25, True), fill=(18, 38, 56))
    d.rounded_rectangle((100, 405, 620, 520), radius=18, fill=WHITE, outline=(216, 224, 233), width=2)
    d.text((125, 432), "I have fever and cough with body pain", font=fnt(22), fill=(78, 92, 110))
    d.rounded_rectangle((100, 548, 280, 594), radius=16, fill=(38, 122, 183))
    d.text((144, 559), "Check Symptoms", font=fnt(18, True), fill=WHITE)
    d.rounded_rectangle((300, 548, 420, 594), radius=16, fill=WHITE, outline=(38, 122, 183), width=2)
    d.text((321, 559), "Clear", font=fnt(18, True), fill=(38, 122, 183))
    d.text((100, 634), "Matched Symptoms", font=fnt(23, True), fill=(18, 38, 56))
    tags = [("fever", (239, 133, 53)), ("cough", (24, 107, 118)), ("body pain", (71, 171, 151))]
    tx = 100
    for tag, color in tags:
        width = 90 + len(tag) * 10
        d.rounded_rectangle((tx, 678, tx + width, 724), radius=18, fill=color)
        d.text((tx + 18, 688), tag, font=fnt(18, True), fill=WHITE)
        tx += width + 16
    d.text((100, 772), "Vague input guard keeps random text from returning unsafe results.", font=fnt(18), fill=(108, 123, 140))

    shadow_card(d, (730, 330, 1530, 860), fill=WHITE)
    d.text((760, 356), "Recommended Medicines", font=fnt(25, True), fill=(18, 38, 56))
    recs = [
        ("Paracetamol", "Fever and pain relief", "Match: 3 symptoms", (239, 133, 53)),
        ("Cough Syrup", "Useful for cough support", "Match: 2 symptoms", (24, 107, 118)),
        ("Vitamin C", "Supportive cold medicine", "Match: 1 symptom", (71, 171, 151)),
    ]
    for i, (name, desc, match, c) in enumerate(recs):
        y = 408 + i * 122
        d.rounded_rectangle((760, y, 1490, y + 96), radius=18, fill=(248, 250, 253), outline=(227, 233, 241), width=1)
        d.ellipse((785, y + 20, 829, y + 64), fill=c)
        d.text((852, y + 16), name, font=fnt(24, True), fill=(18, 38, 56))
        d.text((852, y + 50), desc, font=fnt(17), fill=(96, 110, 126))
        d.rounded_rectangle((1265, y + 22, 1460, y + 60), radius=14, fill=WHITE, outline=c, width=2)
        d.text((1285, y + 29), match, font=fnt(16, True), fill=c)
    d.rounded_rectangle((760, 790, 1490, 840), radius=16, fill=(18, 38, 56))
    d.text((788, 801), "Confidence Score: 82%", font=fnt(20, True), fill=WHITE)
    d.rounded_rectangle((1070, 805, 1450, 825), radius=10, fill=(57, 74, 95))
    d.rounded_rectangle((1070, 805, 1383, 825), radius=10, fill=(71, 171, 151))
    d.text((760, 862), "Live inventory ranking + symptom matching + confidence display", font=fnt(16), fill=(108, 123, 140))

    d.rectangle((0, 962, 1600, 1000), fill=(29, 38, 54))
    d.text((60, 972), "AcuMeds • AI recommendation preview", font=fnt(16), fill=(242, 246, 250))
    img.save(path)


def proposed_solution_preview(path):
    img = Image.new("RGBA", (1600, 1000), (245, 248, 252, 255))
    d = ImageDraw.Draw(img, "RGBA")

    # Soft decorative background shapes to keep the page feeling light.
    d.ellipse((-260, -180, 470, 470), fill=(38, 122, 183, 72))
    d.ellipse((1180, -170, 1900, 420), fill=(71, 171, 151, 70))
    d.ellipse((1240, 520, 1830, 1100), fill=(239, 133, 53, 88))
    d.ellipse((-140, 660, 360, 1180), fill=(24, 107, 118, 28))

    panel = (70, 92, 1530, 800)
    shadow_card(d, panel, fill=WHITE, outline=(223, 230, 239), shadow=(0, 0, 0, 34), offset=(11, 14), radius=34)

    # Header row.
    d.rounded_rectangle((110, 136, 414, 204), radius=22, fill=(239, 133, 53))
    d.text((136, 156), "Proposed Solution", font=fnt(26, True), fill=WHITE)
    d.text((418, 145), "AcuMeds", font=fnt(34, True), fill=(18, 38, 56))
    d.text((418, 192), "AI-Powered Online Pharmacy System", font=fnt(24), fill=(90, 106, 125))

    for i, c in enumerate([(38, 122, 183), (71, 171, 151), (239, 133, 53)]):
        d.ellipse((1260 + i * 46, 145, 1278 + i * 46, 163), fill=c)

    # Left-side story block.
    title_font = fnt(56, True)
    subtitle_font = fnt(17)
    draw_title = "One Platform\nfor Pharmacy Flow"
    d.multiline_text((110, 292), draw_title, font=title_font, fill=(18, 38, 56), spacing=3)
    d.multiline_text(
        (112, 410),
        "Customer, admin, AI, and delivery stay connected.",
        font=subtitle_font,
        fill=(94, 110, 128),
        spacing=6,
    )

    # Central workflow card.
    center_box = (615, 222, 985, 590)
    shadow_card(d, center_box, fill=(247, 250, 254, 255), outline=(213, 222, 233), shadow=(0, 0, 0, 28), offset=(8, 10), radius=30)
    d.ellipse((730, 248, 870, 388), fill=(18, 38, 56))
    d.rectangle((782, 274, 818, 362), fill=WHITE)
    d.rectangle((752, 304, 848, 332), fill=WHITE)
    d.text((652, 430), "Unified pharmacy workflow", font=fnt(28, True), fill=(18, 38, 56))
    d.text((664, 480), "Browse -> Recommend -> Order -> Deliver", font=fnt(18), fill=(95, 109, 126))
    d.rounded_rectangle((684, 520, 914, 560), radius=14, fill=(238, 244, 250), outline=(217, 226, 235), width=1)
    d.text((713, 529), "One shared flow for all roles", font=fnt(16, True), fill=(71, 96, 121))

    # Support cards and connecting lines.
    connector_color = (197, 206, 216, 220)
    connector_width = 7
    lines = [
        ((555, 512), (615, 385)),
        ((555, 686), (615, 550)),
        ((985, 385), (1110, 512)),
        ((985, 550), (1110, 686)),
    ]
    for start, end in lines:
        d.line([start, end], fill=connector_color, width=connector_width)

    def support_card(x0, y0, width, height, accent, title, bullets):
        shadow_card(d, (x0, y0, x0 + width, y0 + height), fill=WHITE, outline=(222, 230, 239), shadow=(0, 0, 0, 20), offset=(5, 6), radius=24)
        d.ellipse((x0 + 18, y0 + 18, x0 + 66, y0 + 66), fill=accent)
        d.text((x0 + 84, y0 + 22), title, font=fnt(25, True), fill=(18, 38, 56))
        bullet_font = fnt(15)
        bullet_y = y0 + 68
        for bullet in bullets:
            d.text((x0 + 24, bullet_y), "• " + bullet, font=bullet_font, fill=(92, 108, 126))
            bullet_y += 28

    support_card(
        110,
        470,
        410,
        150,
        (38, 122, 183),
        "Customer Flow",
        ["Browse medicines", "Add to cart and checkout", "Upload prescription when needed"],
    )
    support_card(
        110,
        642,
        410,
        150,
        (239, 133, 53),
        "Admin Flow",
        ["Manage medicines and stock", "Verify prescriptions and shopkeepers", "Track orders and reports"],
    )
    support_card(
        1110,
        470,
        410,
        150,
        (71, 171, 151),
        "AI Recommendation",
        ["Read symptom text", "Match medicine suggestions", "Show reasons and confidence"],
    )
    support_card(
        1110,
        642,
        410,
        150,
        (24, 107, 118),
        "Delivery Flow",
        ["Assign order to agent", "Update delivery status", "Confirm successful delivery"],
    )

    # Footer goal bar.
    d.rounded_rectangle((98, 860, 1502, 922), radius=26, fill=(18, 38, 56))
    d.text((126, 879), "Goal: improve convenience, accuracy, and traceability across pharmacy operations.", font=fnt(17), fill=WHITE)

    img.save(path)


def gallery_preview(path):
    img = Image.new("RGBA", (1600, 1000), (245, 248, 252, 255))
    d = ImageDraw.Draw(img, "RGBA")

    d.text((60, 48), "Project Gallery", font=fnt(40, True), fill=(18, 38, 56))
    d.text((60, 100), "Website previews and medicine assets used in the AcuMeds project", font=fnt(20), fill=(92, 106, 122))

    tile_positions = [
        (60, 170, 460, 305),
        (570, 170, 460, 305),
        (1080, 170, 460, 305),
        (60, 535, 460, 305),
        (570, 535, 460, 305),
        (1080, 535, 460, 305),
    ]

    for (src, label, accent), (x, y, w, h) in zip(MEDICINE_SOURCE_IMAGES, tile_positions):
        shadow_card(d, (x, y, x + w, y + h), fill=WHITE)
        d.rounded_rectangle((x, y, x + w, y + 76), radius=20, fill=accent)
        d.text((x + 20, y + 22), label, font=fnt(24, True), fill=WHITE)

        inner = (x + 18, y + 88, x + w - 18, y + h - 18)
        if src.exists():
            try:
                source_img = Image.open(src).convert("RGBA")
                thumb = ImageOps.contain(source_img, (inner[2] - inner[0], inner[3] - inner[1]))
                px = inner[0] + ((inner[2] - inner[0]) - thumb.width) // 2
                py = inner[1] + ((inner[3] - inner[1]) - thumb.height) // 2
                img.alpha_composite(thumb, (px, py))
            except Exception:
                d.rounded_rectangle(inner, radius=16, fill=(243, 246, 250), outline=accent, width=2)
                d.text((inner[0] + 24, inner[1] + 86), "Image unavailable", font=fnt(18, True), fill=(92, 106, 122))
        else:
            d.rounded_rectangle(inner, radius=16, fill=(243, 246, 250), outline=accent, width=2)
            d.text((inner[0] + 24, inner[1] + 86), "Image unavailable", font=fnt(18, True), fill=(92, 106, 122))

    d.rounded_rectangle((60, 870, 1540, 934), radius=18, fill=(18, 38, 56))
    d.text((84, 887), "This collage is used on the gallery slide to make the presentation more visual and easier to present.", font=fnt(17), fill=WHITE)
    img.save(path)


home_preview(HOME_PNG)
ai_preview(AI_PNG)
proposed_solution_preview(PROPOSED_PNG)
gallery_preview(GALLERY_PNG)

# ---------- PPT helpers ----------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "AcuMeds Project Presentation"
prs.core_properties.subject = "Expanded engineering PPT"
prs.core_properties.author = "GitHub Copilot"


def set_bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def shp(slide, left, top, width, height, fill=CARD, line=LINE, radius=True):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def txt(slide, left, top, width, height, content, *, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.clear()
    lines = content if isinstance(content, list) else str(content).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(4)
        if p.runs:
            r = p.runs[0]
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
    return tb


def header(slide, title, page_no, agenda_slide=None):
    set_bg(slide)
    shp(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.82), fill=NAVY, line=NAVY, radius=False)
    shp(slide, Inches(0), Inches(0.82), Inches(13.333), Inches(0.08), fill=ORANGE, line=ORANGE, radius=False)
    txt(slide, Inches(0.55), Inches(0.18), Inches(8.9), Inches(0.32), title, size=24, color=WHITE, bold=True)
    if agenda_slide is not None:
        menu_btn = shp(slide, Inches(9.48), Inches(0.16), Inches(0.95), Inches(0.34), fill=WHITE, line=WHITE)
        txt(slide, Inches(9.58), Inches(0.21), Inches(0.75), Inches(0.16), "Agenda", size=10, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        menu_btn.click_action.target_slide = agenda_slide
    badge = shp(slide, Inches(10.55), Inches(0.16), Inches(2.05), Inches(0.34), fill=TEAL, line=TEAL)
    txt(slide, Inches(10.72), Inches(0.21), Inches(1.7), Inches(0.16), "AcuMeds", size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    if agenda_slide is not None:
        badge.click_action.target_slide = agenda_slide
    txt(slide, Inches(12.35), Inches(7.0), Inches(0.55), Inches(0.16), str(page_no), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def card(slide, left, top, width, height, title, body, accent=ORANGE, title_size=18, body_size=13, centered=False):
    base = shp(slide, left, top, width, height, fill=CARD, line=LINE, radius=True)
    shp(slide, left, top, width, Inches(0.08), fill=accent, line=accent, radius=False)
    txt(slide, left + Inches(0.18), top + Inches(0.14), width - Inches(0.36), Inches(0.34), title, size=title_size, color=TEXT, bold=True, align=PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT)
    if body:
        txt(slide, left + Inches(0.18), top + Inches(0.55), width - Inches(0.36), height - Inches(0.7), body, size=body_size, color=MUTED, align=PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT)
    return base


def flow(slide, left, top, width, height, title, subtitle, fill):
    shp(slide, left, top, width, height, fill=fill, line=fill, radius=True)
    shp(slide, left, top, width, Inches(0.08), fill=WHITE, line=WHITE, radius=False)
    txt(slide, left + Inches(0.14), top + Inches(0.16), width - Inches(0.28), Inches(0.26), title, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    txt(slide, left + Inches(0.12), top + Inches(0.54), width - Inches(0.24), height - Inches(0.58), subtitle, size=11, color=WHITE, align=PP_ALIGN.CENTER)


def arrow(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = LINE
    shape.line.color.rgb = LINE
    return shape


# Title slide
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
shp(slide, Inches(9.0), Inches(-0.1), Inches(4.4), Inches(4.4), fill=ORANGE, line=ORANGE)
shp(slide, Inches(10.1), Inches(1.0), Inches(2.9), Inches(2.9), fill=BLUE, line=BLUE)
shp(slide, Inches(0.8), Inches(0.95), Inches(0.18), Inches(5.15), fill=ORANGE, line=ORANGE, radius=False)
txt(slide, Inches(1.1), Inches(1.46), Inches(8.4), Inches(0.95), "AcuMeds", size=46, color=WHITE, bold=True)
txt(slide, Inches(1.15), Inches(2.45), Inches(7.8), Inches(0.5), "AI-Powered Online Pharmacy System", size=22, color=WHITE)
txt(slide, Inches(1.15), Inches(3.05), Inches(7.8), Inches(0.6), "Expanded Engineering Project Presentation", size=17, color=RGBColor(214, 223, 231))
txt(slide, Inches(1.15), Inches(4.08), Inches(7.95), Inches(0.82), "This presentation covers the full project flow: problem statement, proposed solution, system architecture, AI recommendation module, user roles, website screenshots, workflow, technologies, result, benefits, differentiation, features, scope, and conclusion.", size=13, color=RGBColor(214, 223, 231))
card(slide, Inches(1.15), Inches(5.25), Inches(2.05), Inches(0.92), "Platform", "Flask + SQLite", accent=TEAL, title_size=16, body_size=11, centered=True)
card(slide, Inches(3.4), Inches(5.25), Inches(2.1), Inches(0.92), "Scope", "Academic + practical", accent=BLUE, title_size=16, body_size=11, centered=True)
card(slide, Inches(5.7), Inches(5.25), Inches(2.1), Inches(0.92), "Focus", "Healthtech workflow", accent=MINT, title_size=16, body_size=11, centered=True)

title_slide = slide

agenda_slide = prs.slides.add_slide(prs.slide_layouts[6])
header(agenda_slide, "Agenda", 2, agenda_slide=agenda_slide)
txt(agenda_slide, Inches(0.75), Inches(1.18), Inches(6.2), Inches(0.36), "Tap a card to jump to that section", size=13, color=MUTED)
txt(agenda_slide, Inches(0.75), Inches(1.48), Inches(6.0), Inches(0.35), "Interactive topic map for the full presentation", size=18, color=TEXT, bold=True)

agenda_items = []
agenda_specs = [
    ("Introduction & Problem", "Start with the need and context", ORANGE, "Introduction", 0.75, 2.0),
    ("Proposed Solution", "What AcuMeds solves", BLUE, "Proposed Solution", 3.85, 2.0),
    ("System Architecture", "How the app is structured", TEAL, "System Architecture", 0.75, 3.1),
    ("AI Recommendation", "Symptom-based medicine support", MINT, "AI Recommendation Module", 3.85, 3.1),
    ("User Roles", "Customer, admin, shopkeeper, delivery", ORANGE, "User Roles", 0.75, 4.2),
    ("Website & Gallery", "Screenshots and medicine images", BLUE, "Website Screenshot", 3.85, 4.2),
    ("Workflow & Tech", "Process and tools used", MINT, "Workflow", 0.75, 5.3),
    ("Benefits & Scope", "Why it matters and future work", TEAL, "Scope", 3.85, 5.3),
]
for title_text, subtitle, accent, target_title, left_pos, top_pos in agenda_specs:
    agenda_shape = card(
        agenda_slide,
        Inches(left_pos),
        Inches(top_pos),
        Inches(2.85),
        Inches(0.82),
        title_text,
        subtitle,
        accent=accent,
        title_size=15,
        body_size=10,
        centered=True,
    )
    agenda_items.append((agenda_shape, target_title))

shp(agenda_slide, Inches(7.15), Inches(1.28), Inches(5.45), Inches(2.35), fill=WHITE, line=LINE, radius=True)
shp(agenda_slide, Inches(7.15), Inches(3.82), Inches(5.45), Inches(2.35), fill=WHITE, line=LINE, radius=True)
agenda_slide.shapes.add_picture(str(HOME_PNG), Inches(7.28), Inches(1.41), width=Inches(5.19), height=Inches(2.1))
agenda_slide.shapes.add_picture(str(AI_PNG), Inches(7.28), Inches(3.95), width=Inches(5.19), height=Inches(2.1))
txt(agenda_slide, Inches(7.32), Inches(1.31), Inches(2.5), Inches(0.22), "Homepage UI", size=10, color=WHITE, bold=True)
txt(agenda_slide, Inches(7.32), Inches(3.85), Inches(3.0), Inches(0.22), "AI Recommendation UI", size=10, color=WHITE, bold=True)

agenda_button = shp(title_slide, Inches(8.55), Inches(5.25), Inches(2.15), Inches(0.62), fill=ORANGE, line=ORANGE, radius=True)
txt(title_slide, Inches(8.72), Inches(5.41), Inches(1.8), Inches(0.18), "Open Agenda", size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
agenda_button.click_action.target_slide = agenda_slide

# Content slides
slides = [
    ("Introduction", ORANGE, [
        "AcuMeds is an AI-powered online pharmacy and medicine management system.",
        "It allows users to browse medicines, search products, upload prescriptions, and place orders online.",
        "The platform also supports admins, shopkeepers, and delivery agents in one modular application.",
        "The objective is to make medicine access simple, organized, and traceable for daily pharmacy use."
    ]),
    ("Problem Statement", BLUE, [
        "Users often do not know which medicine matches a symptom such as fever, cold, cough, headache, or allergy.",
        "Traditional pharmacy processes are slow because stock handling, prescriptions, billing, and delivery may be managed separately.",
        "Order status tracking is weak when the workflow is not centralized.",
        "Pharmacies need one secure system that supports customers and staff together."
    ]),
    ("Proposed Solution", MINT, None),
    ("System Architecture", TEAL, None),
    ("AI Recommendation Module", ORANGE, None),
    ("User Roles", BLUE, None),
    ("Website Screenshot", MINT, None),
    ("Project Gallery", TEAL, None),
    ("Workflow", TEAL, None),
    ("Technologies", ORANGE, [
        "Backend: Python, Flask, Flask-Login, SQLAlchemy",
        "Database: SQLite with structured relationships",
        "Frontend: HTML5, CSS3, Bootstrap 5.3, JavaScript",
        "AI / logic: TF-IDF-assisted symptom matching, alias handling, and confidence scoring",
        "Supporting tools: QR code generation, PDF / HTML reports, static assets, and modular templates"
    ]),
    ("Features", BLUE, None),
    ("Result", MINT, [
        "The project successfully unifies medicine browsing, recommendation, ordering, and delivery tracking.",
        "It provides a complete database-backed workflow for all major pharmacy operations.",
        "Users receive useful symptom-based suggestions while admins keep full control over inventory and orders.",
        "The result is a practical, real-world style engineering solution for pharmacy management."
    ]),
    ("Benefits", TEAL, None),
    ("How Different From Other Systems", ORANGE, None),
    ("Scope", BLUE, [
        "Future enhancements can include payment gateway integration, notifications, and SMS alerts.",
        "The recommendation engine can be improved with stronger medical validation and interaction checks.",
        "Mobile app support, analytics, multilingual support, and doctor consultation can be added later.",
        "The modular design already gives a strong base for future expansion."
    ]),
    ("Conclusion", MINT, [
        "AcuMeds is a complete engineering project that combines pharmacy commerce, AI support, and delivery workflows.",
        "It improves convenience, reduces manual work, and gives better control to both users and administrators.",
        "The modular Flask architecture makes it maintainable and extensible for future upgrades.",
        "Thank you."
    ]),
]

section_slides = {}

for idx, (title, accent, body_lines) in enumerate(slides, start=2):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    section_slides[title] = slide

    section_image = SECTION_IMAGE_PATHS.get(title)
    if section_image and section_image.exists():
        slide.shapes.add_picture(str(section_image), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
        continue

    header(slide, title, idx, agenda_slide=agenda_slide)

    if title == "Proposed Solution":
        card(slide, Inches(0.65), Inches(1.35), Inches(3.95), Inches(4.95), "Customer Flow", [
            "Browse medicines",
            "Search and filter products",
            "Add items to cart",
            "Checkout with address and notes"
        ], accent=ORANGE, title_size=18, body_size=13, centered=True)
        card(slide, Inches(4.65), Inches(1.35), Inches(3.95), Inches(4.95), "Admin Flow", [
            "Manage medicines and stock",
            "Verify prescriptions and shopkeepers",
            "Track orders and delivery",
            "View dashboard reports"
        ], accent=BLUE, title_size=18, body_size=13, centered=True)
        card(slide, Inches(8.65), Inches(1.35), Inches(3.95), Inches(4.95), "AI Support", [
            "Accept symptom text",
            "Match medicine suggestions",
            "Show reasons and confidence",
            "Store recommendation history"
        ], accent=MINT, title_size=18, body_size=13, centered=True)

    elif title == "System Architecture":
        x = Inches(0.42)
        y = Inches(2.0)
        widths = [Inches(1.95), Inches(2.18), Inches(2.15), Inches(2.1), Inches(2.05)]
        layers = [
            ("User Layer", "Customer\nAdmin\nShopkeeper\nDelivery Agent", TEAL),
            ("Application Layer", "Flask blueprints\nAuth\nCart\nOrders\nAdmin", BLUE),
            ("AI Layer", "Symptom preprocessing\nMatching\nRanking\nConfidence", ORANGE),
            ("Data Layer", "SQLAlchemy models\nSQLite database", MINT),
            ("Output Layer", "Dashboards\nHistory\nTracking\nReports", TEAL),
        ]
        for i, ((title_text, sub, fill), width) in enumerate(zip(layers, widths)):
            flow(slide, x, y, width, Inches(1.7), title_text, sub, fill)
            x = x + width
            if i < len(layers) - 1:
                arrow(slide, x + Inches(0.03), y + Inches(0.62), Inches(0.22), Inches(0.36))
                x = x + Inches(0.3)
        card(slide, Inches(0.75), Inches(4.55), Inches(12.0), Inches(1.3), "Architecture Note", [
            "The layered design separates user interface, business logic, AI logic, and database storage, making the system easier to maintain and extend."
        ], accent=BLUE, title_size=16, body_size=12, centered=True)

    elif title == "AI Recommendation Module":
        card(slide, Inches(0.65), Inches(1.35), Inches(12.1), Inches(1.45), "How It Works", [
            "User enters symptoms -> text is normalized -> supported symptoms and aliases are matched -> live inventory is scored -> ranked medicines and confidence are returned."
        ], accent=ORANGE, title_size=19, body_size=13, centered=True)
        card(slide, Inches(0.65), Inches(3.0), Inches(5.95), Inches(3.3), "What the module does", [
            "Recognizes natural symptom wording such as fever, cough, cold, headache, body pain, or allergy.",
            "Ranks medicines using the current inventory, so suggestions come from real available items.",
            "Provides a confidence score and stores recommendation history for later reference."
        ], accent=BLUE, title_size=18, body_size=13)
        card(slide, Inches(6.8), Inches(3.0), Inches(5.95), Inches(3.3), "Why it is useful", [
            "Reduces confusion for users who do not know medicine names.",
            "Improves explanation by showing matched symptoms and reasons.",
            "Avoids unsafe random outputs by ignoring vague or unsupported inputs."
        ], accent=MINT, title_size=18, body_size=13)

    elif title == "User Roles":
        roles = [
            ("Customer", ["Register, login, browse medicines", "Add to cart and place orders", "Upload prescription if needed"], ORANGE),
            ("Admin", ["Manage medicines and stock", "Approve shopkeepers and monitor users", "Track orders, reports, and settings"], BLUE),
            ("Shopkeeper", ["Add and update medicines", "Handle inventory and fulfillment", "Review shop-specific orders"], MINT),
            ("Delivery Agent", ["Pick assigned orders", "Update delivery status", "Confirm delivery and proof"], TEAL),
        ]
        xs = [0.55, 3.72, 6.89, 10.06]
        for (title_text, lines, fill), x in zip(roles, xs):
            card(slide, Inches(x), Inches(1.55), Inches(2.75), Inches(4.8), title_text, ["- " + line for line in lines], accent=fill, title_size=18, body_size=12, centered=True)
        card(slide, Inches(0.75), Inches(6.55), Inches(12.0), Inches(0.45), "Role-based access keeps user actions separated and makes the platform more secure and realistic.", None, accent=ORANGE, title_size=12, body_size=10, centered=True)

    elif title == "Website Screenshot":
        txt(slide, Inches(0.75), Inches(1.25), Inches(12.0), Inches(0.32), "Homepage and recommendation preview from the project UI", size=12, color=MUTED, align=PP_ALIGN.CENTER)
        shp(slide, Inches(0.55), Inches(1.55), Inches(7.5), Inches(5.45), fill=WHITE, line=LINE, radius=True)
        shp(slide, Inches(8.25), Inches(1.55), Inches(4.5), Inches(5.45), fill=WHITE, line=LINE, radius=True)
        slide.shapes.add_picture(str(HOME_PNG), Inches(0.72), Inches(1.72), width=Inches(7.15), height=Inches(5.0))
        slide.shapes.add_picture(str(AI_PNG), Inches(8.42), Inches(1.72), width=Inches(4.15), height=Inches(5.0))
        txt(slide, Inches(0.72), Inches(1.58), Inches(7.0), Inches(0.25), "Homepage screenshot", size=11, color=WHITE, bold=True)
        txt(slide, Inches(8.42), Inches(1.58), Inches(4.0), Inches(0.25), "AI recommendation screenshot", size=11, color=WHITE, bold=True)
        card(slide, Inches(0.75), Inches(6.85), Inches(12.0), Inches(0.28), "The UI shows the pharmacy homepage, featured medicines, and the symptom-based AI recommendation flow.", None, accent=TEAL, title_size=12, body_size=10, centered=True)

    elif title == "Project Gallery":
        txt(slide, Inches(0.75), Inches(1.22), Inches(12.0), Inches(0.3), "Medicine images and UI previews from the project", size=12, color=MUTED, align=PP_ALIGN.CENTER)
        shp(slide, Inches(0.55), Inches(1.55), Inches(12.25), Inches(4.9), fill=WHITE, line=LINE, radius=True)
        slide.shapes.add_picture(str(GALLERY_PNG), Inches(0.8), Inches(1.82), width=Inches(11.75), height=Inches(4.4))
        card(slide, Inches(0.75), Inches(6.6), Inches(12.0), Inches(0.45), "The gallery brings together the homepage, AI module, and medicine assets so the deck looks more visual and presentation-friendly.", None, accent=MINT, title_size=12, body_size=10, centered=True)

    elif title == "Workflow":
        steps = [
            ("1", "Register / Login", TEAL),
            ("2", "Browse Medicines", BLUE),
            ("3", "AI Check Symptoms", ORANGE),
            ("4", "Add to Cart", MINT),
            ("5", "Checkout & Pay", TEAL),
            ("6", "Delivery Tracking", BLUE),
        ]
        x = Inches(0.4)
        y = Inches(2.0)
        for i, (num, label, fill) in enumerate(steps):
            shp(slide, x, y, Inches(2.0), Inches(1.6), fill=fill, line=fill, radius=True)
            txt(slide, x + Inches(0.12), y + Inches(0.16), Inches(1.76), Inches(0.28), num, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
            txt(slide, x + Inches(0.1), y + Inches(0.55), Inches(1.8), Inches(0.32), label, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
            x = x + Inches(2.0)
            if i < len(steps) - 1:
                arrow(slide, x + Inches(0.02), y + Inches(0.63), Inches(0.18), Inches(0.28))
                x = x + Inches(0.22)
        card(slide, Inches(0.75), Inches(4.15), Inches(12.0), Inches(1.6), "Workflow Summary", [
            "The workflow starts from user login, moves through product discovery and AI guidance, and ends with payment, order fulfillment, and delivery tracking."
        ], accent=ORANGE, title_size=17, body_size=12, centered=True)

    elif title == "Technologies":
        tech_cards = [
            ("Backend", ["Python", "Flask", "Flask-Login"], TEAL),
            ("Database", ["SQLite", "SQLAlchemy", "Relational models"], BLUE),
            ("Frontend", ["HTML5", "CSS3", "Bootstrap 5.3", "JavaScript"], ORANGE),
            ("AI / Logic", ["TF-IDF assisted matching", "Aliases", "Confidence score"], MINT),
            ("Tools", ["QR code generation", "PDF / HTML reports", "Static assets"], TEAL),
        ]
        positions = [(0.55, 1.35, 3.85, 1.9), (4.55, 1.35, 3.85, 1.9), (8.55, 1.35, 4.2, 1.9), (0.55, 3.55, 5.9, 2.0), (6.65, 3.55, 6.1, 2.0)]
        for (title_text, lines, fill), (x, y, w, h) in zip(tech_cards, positions):
            card(slide, Inches(x), Inches(y), Inches(w), Inches(h), title_text, ["- " + line for line in lines], accent=fill, title_size=18, body_size=12, centered=True)

    elif title == "Features":
        card(slide, Inches(0.6), Inches(1.35), Inches(3.9), Inches(4.95), "User Features", [
            "Registration and secure login",
            "Medicine search and filtering",
            "Cart, checkout, and order history",
            "Prescription upload and tracking"
        ], accent=ORANGE, title_size=18, body_size=13, centered=True)
        card(slide, Inches(4.67), Inches(1.35), Inches(3.9), Inches(4.95), "Admin Features", [
            "Dashboard and reporting",
            "Medicine CRUD operations",
            "Order monitoring and status updates",
            "User and verification control"
        ], accent=BLUE, title_size=18, body_size=13, centered=True)
        card(slide, Inches(8.74), Inches(1.35), Inches(3.9), Inches(4.95), "AI Features", [
            "Symptom-based medicine suggestions",
            "Alias and keyword matching",
            "Confidence scoring and explanations",
            "Recommendation history storage"
        ], accent=MINT, title_size=18, body_size=13, centered=True)

    elif title == "Benefits":
        card(slide, Inches(0.6), Inches(1.35), Inches(3.9), Inches(4.95), "For Users", [
            "Faster medicine discovery",
            "Simple shopping experience",
            "Helpful symptom guidance",
            "Clear delivery and payment tracking"
        ], accent=TEAL, title_size=18, body_size=13, centered=True)
        card(slide, Inches(4.67), Inches(1.35), Inches(3.9), Inches(4.95), "For Pharmacy", [
            "Centralized stock control",
            "Reduced manual work",
            "Better order handling",
            "Improved operational visibility"
        ], accent=ORANGE, title_size=18, body_size=13, centered=True)
        card(slide, Inches(8.74), Inches(1.35), Inches(3.9), Inches(4.95), "For Management", [
            "Role-based access control",
            "Traceable records",
            "Better reporting and oversight",
            "Foundation for future scaling"
        ], accent=BLUE, title_size=18, body_size=13, centered=True)

    elif title == "How Different From Other Systems":
        card(slide, Inches(0.6), Inches(1.35), Inches(3.95), Inches(4.95), "Traditional Systems", [
            "Mostly catalog-based",
            "Manual symptom search",
            "Limited role support",
            "Weak workflow tracking"
        ], accent=BLUE, title_size=18, body_size=13, centered=True)
        card(slide, Inches(4.67), Inches(1.35), Inches(3.95), Inches(4.95), "AcuMeds Advantage", [
            "AI symptom recommendation",
            "Prescription and delivery workflow",
            "Admin / shopkeeper / delivery roles",
            "Recommendation history and reports"
        ], accent=ORANGE, title_size=18, body_size=13, centered=True)
        card(slide, Inches(8.74), Inches(1.35), Inches(3.9), Inches(4.95), "Why It Stands Out", [
            "Live inventory ranked suggestions",
            "Explainable output with reasons",
            "Multi-role pharmacy workflow",
            "Closer to a real engineering solution"
        ], accent=MINT, title_size=18, body_size=13, centered=True)

    elif title == "Scope":
        shp(slide, Inches(0.55), Inches(1.12), Inches(12.23), Inches(5.95), fill=WHITE, line=LINE, radius=True)
        shp(slide, Inches(0.8), Inches(1.2), Inches(11.73), Inches(0.74), fill=NAVY, line=NAVY, radius=True)
        txt(slide, Inches(1.0), Inches(1.34), Inches(3.4), Inches(0.26), "Scope", size=32, color=WHITE, bold=True)
        txt(slide, Inches(1.0), Inches(1.6), Inches(5.0), Inches(0.18), "Future enhancements and expansion areas for AcuMeds.", size=11, color=RGBColor(216, 223, 232))

        scope_chips = [
            ("Enhance", BLUE, 8.55, 0.74),
            ("Validate", ORANGE, 9.36, 0.77),
            ("Expand", MINT, 10.19, 0.73),
            ("Mobile", TEAL, 11.0, 0.72),
        ]
        for label, accent, left_pos, width in scope_chips:
            shp(slide, Inches(left_pos), Inches(1.28), Inches(width), Inches(0.33), fill=accent, line=accent, radius=True)
            txt(slide, Inches(left_pos), Inches(1.33), Inches(width), Inches(0.15), label, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

        txt(slide, Inches(0.95), Inches(2.0), Inches(8.6), Inches(0.24), "Scope is focused on improving safety, accessibility, and the system base for future growth.", size=14, color=MUTED)

        def scope_panel(left, top, width, height, accent, number, panel_title, subtitle, bullets):
            shp(slide, left, top, width, height, fill=WHITE, line=LINE, radius=True)
            shp(slide, left, top, width, Inches(0.08), fill=accent, line=accent, radius=False)
            badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, left + Inches(0.16), top + Inches(0.18), Inches(0.42), Inches(0.42))
            badge.fill.solid()
            badge.fill.fore_color.rgb = accent
            badge.line.color.rgb = accent
            txt(slide, left + Inches(0.16), top + Inches(0.19), Inches(0.42), Inches(0.22), str(number), size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
            txt(slide, left + Inches(0.72), top + Inches(0.22), width - Inches(0.92), Inches(0.24), panel_title, size=18, color=TEXT, bold=True)
            txt(slide, left + Inches(0.72), top + Inches(0.48), width - Inches(0.92), Inches(0.18), subtitle, size=10, color=accent)
            bullet_lines = ["• " + line for line in bullets]
            txt(slide, left + Inches(0.16), top + Inches(0.83), width - Inches(0.32), height - Inches(0.92), bullet_lines, size=11, color=MUTED)

        scope_panel(Inches(0.8), Inches(2.25), Inches(4.15), Inches(1.55), BLUE, 1, "Payments & Alerts", "Make the experience more complete", [
            "Payment gateway integration",
            "Push notifications for updates",
            "SMS alerts for order status",
        ])
        scope_panel(Inches(8.37), Inches(2.25), Inches(4.15), Inches(1.55), ORANGE, 2, "Safety & Validation", "Improve reliability of recommendations", [
            "Stronger medical validation",
            "Interaction checks for medicines",
            "Safer recommendation filters",
        ])
        scope_panel(Inches(0.8), Inches(4.37), Inches(4.15), Inches(1.55), TEAL, 3, "Mobile & Analytics", "Increase reach and visibility", [
            "Mobile app support",
            "Analytics dashboards",
            "Multilingual support",
        ])
        scope_panel(Inches(8.37), Inches(4.37), Inches(4.15), Inches(1.55), MINT, 4, "Consultation & Scale", "Keep the app ready for growth", [
            "Doctor consultation flow",
            "Modular design for expansion",
            "Easy future feature upgrades",
        ])

        center_left = Inches(5.12)
        center_top = Inches(2.1)
        center_width = Inches(3.05)
        center_height = Inches(3.88)
        shp(slide, center_left, center_top, center_width, center_height, fill=RGBColor(248, 251, 254), line=LINE, radius=True)
        center_circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, center_left + Inches(0.95), center_top + Inches(0.22), Inches(1.18), Inches(1.18))
        center_circle.fill.solid()
        center_circle.fill.fore_color.rgb = NAVY
        center_circle.line.color.rgb = NAVY
        txt(slide, center_left + Inches(1.02), center_top + Inches(0.34), Inches(1.0), Inches(0.45), "+", size=42, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        txt(slide, center_left + Inches(0.42), center_top + Inches(1.48), Inches(2.2), Inches(0.32), "AcuMeds Scope", size=23, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        txt(slide, center_left + Inches(0.24), center_top + Inches(1.88), Inches(2.58), Inches(0.58), "A strong base is already in place\nfor safety, scale, and accessibility upgrades", size=13, color=MUTED, align=PP_ALIGN.CENTER)

        small_tags = [
            ("Safety", BLUE, 5.25, 4.05, 0.98),
            ("Scale", MINT, 6.32, 4.05, 0.88),
            ("Accessibility", TEAL, 5.25, 4.47, 1.22),
            ("Expand", ORANGE, 6.56, 4.47, 0.86),
        ]
        for label, accent, left_pos, top_pos, width in small_tags:
            shp(slide, Inches(left_pos), Inches(top_pos), Inches(width), Inches(0.3), fill=RGBColor(238, 244, 250), line=RGBColor(238, 244, 250), radius=True)
            txt(slide, Inches(left_pos), Inches(top_pos) + Inches(0.05), Inches(width), Inches(0.12), label, size=9, color=accent, bold=True, align=PP_ALIGN.CENTER)

        arrow(slide, Inches(4.88), Inches(2.84), Inches(0.42), Inches(0.16))
        arrow(slide, Inches(8.2), Inches(2.84), Inches(0.42), Inches(0.16))
        arrow(slide, Inches(4.88), Inches(4.14), Inches(0.42), Inches(0.16))
        arrow(slide, Inches(8.2), Inches(4.14), Inches(0.42), Inches(0.16))

        shp(slide, Inches(0.8), Inches(6.45), Inches(11.73), Inches(0.56), fill=NAVY, line=NAVY, radius=True)
        txt(slide, Inches(1.0), Inches(6.62), Inches(11.1), Inches(0.16), "This scope keeps the project practical today and extensible for future versions.", size=13, color=WHITE)

    else:
        card(slide, Inches(0.75), Inches(1.35), Inches(12.0), Inches(5.25), title, ["- " + line for line in body_lines], accent=accent, title_size=20, body_size=14)

for agenda_shape, target_title in agenda_items:
    agenda_shape.click_action.target_slide = section_slides[target_title]

# Closing slide
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
shp(slide, Inches(9.0), Inches(-0.1), Inches(4.4), Inches(4.4), fill=ORANGE, line=ORANGE)
shp(slide, Inches(10.1), Inches(1.0), Inches(2.9), Inches(2.9), fill=BLUE, line=BLUE)
shp(slide, Inches(0.8), Inches(0.95), Inches(0.18), Inches(5.15), fill=ORANGE, line=ORANGE, radius=False)
txt(slide, Inches(1.1), Inches(1.75), Inches(8.4), Inches(0.9), "Thank You", size=44, color=WHITE, bold=True)
txt(slide, Inches(1.15), Inches(2.72), Inches(7.9), Inches(0.5), "AcuMeds - AI-Powered Online Pharmacy System", size=20, color=RGBColor(214, 223, 231))
txt(slide, Inches(1.15), Inches(3.4), Inches(8.1), Inches(0.8), "This project shows how one modular web application can combine medicine search, AI guidance, order management, and delivery workflow.", size=15, color=RGBColor(214, 223, 231))
card(slide, Inches(1.15), Inches(4.95), Inches(4.2), Inches(0.95), "Presentation Ready", "For project submission and viva", accent=TEAL, title_size=18, body_size=12, centered=True)
card(slide, Inches(5.6), Inches(4.95), Inches(3.2), Inches(0.95), "Slides", "19 total", accent=BLUE, title_size=18, body_size=12, centered=True)

prs.save(OUT)
print(f"Saved {OUT}")
print(f"Slides: {len(prs.slides)}")
print(f"Images: {HOME_PNG.exists()}, {AI_PNG.exists()}")